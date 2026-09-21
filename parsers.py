"""
parsers.py - Universal document reader for SDOC hackathon attachments.

Handles:
  - .txt  (UTF-8 / fallback latin-1)
  - .pdf  (pypdf text extraction, detection of corrupted or image-only scanned files)
  - .docx (python-docx paragraphs and tables)
  - .xlsx (openpyxl sheet rows and cells)

Also identifies wrong document types (Commercial Invoices, Packing Lists, Certificates of Origin)
and unreadable files (corrupted PDFs, image-only scans without text).
"""
import io
import re
from pathlib import Path

# Headers / markers indicating wrong document type
WRONG_DOC_MARKERS = [
    (r"\bCOMMERCIAL\s+INVOICE\b", "wrong_doc_type"),
    (r"\bPACKING\s+LIST\b", "wrong_doc_type"),
    (r"\bCERTIFICATE\s+OF\s+ORIGIN\b", "wrong_doc_type"),
    (r"\bINVOICE\s+NO\b", "wrong_doc_type"),
    (r"NOT\s+A\s+SHIPPING\s+INSTRUCTION", "wrong_doc_type"),
    (r"NOT\s+A\s+DRAFT\s+BILL\s+OF\s+LADING", "wrong_doc_type"),
    (r"NOT\s+AN\s+SI\s+OR\s+BL", "wrong_doc_type"),
]


def extract_txt(raw_bytes: bytes) -> str:
    """Extract text from plain text bytes."""
    try:
        return raw_bytes.decode("utf-8")
    except UnicodeDecodeError:
        return raw_bytes.decode("latin-1", errors="replace")


def extract_docx(raw_bytes: bytes) -> str:
    """Extract paragraphs and tables from .docx bytes."""
    import docx

    doc = docx.Document(io.BytesIO(raw_bytes))
    lines = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            lines.append(text)
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if row_text:
                lines.append(row_text)
    return "\n".join(lines)


def extract_xlsx(raw_bytes: bytes) -> str:
    """Extract rows and cell contents from .xlsx bytes."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
    lines = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip() != ""]
            if cells:
                lines.append(" : ".join(cells))
    return "\n".join(lines)


def extract_pdf(raw_bytes: bytes) -> tuple[str | None, str | None]:
    """
    Extract text from PDF.
    Returns (text, error_type).
    If corrupted or image-only scanned without text, returns (None, 'unreadable').
    """
    import pypdf

    try:
        reader = pypdf.PdfReader(io.BytesIO(raw_bytes))
        if len(reader.pages) == 0:
            return None, "unreadable"

        pages_text = []
        for page in reader.pages:
            t = page.extract_text()
            if t:
                pages_text.append(t.strip())

        full_text = "\n".join(pages_text).strip()
        # If the PDF has virtually no extractable text, it's an image-only scan
        if len(full_text) < 30:
            return None, "unreadable"

        return full_text, None
    except Exception:
        # Corrupted PDF (e.g. email_511_BL.pdf)
        return None, "unreadable"


def check_wrong_doc_type(text: str) -> str | None:
    """Check if the document text matches a wrong document type."""
    if not text:
        return None
    for pattern, reason in WRONG_DOC_MARKERS:
        if re.search(pattern, text, re.IGNORECASE):
            return reason
    return None


def parse_attachment(file_path: str, raw_bytes: bytes) -> tuple[str | None, str | None]:
    """
    Universal attachment parser.
    Returns:
        (extracted_text, review_reason)
    Where review_reason can be:
        None (parsed cleanly)
        'unreadable' (corrupted, empty, or unparseable)
        'wrong_doc_type' (commercial invoice, packing list, certificate of origin)
    """
    ext = Path(file_path).suffix.lower()

    if not raw_bytes or len(raw_bytes) == 0:
        return None, "unreadable"

    try:
        if ext == ".txt":
            text = extract_txt(raw_bytes)
            reason = check_wrong_doc_type(text)
            if reason:
                return text, reason
            return text, None

        elif ext == ".pdf":
            text, reason = extract_pdf(raw_bytes)
            if reason:
                return None, reason
            doc_reason = check_wrong_doc_type(text)
            if doc_reason:
                return text, doc_reason
            return text, None

        elif ext == ".docx":
            text = extract_docx(raw_bytes)
            reason = check_wrong_doc_type(text)
            if reason:
                return text, reason
            return text, None

        elif ext == ".xlsx":
            text = extract_xlsx(raw_bytes)
            reason = check_wrong_doc_type(text)
            if reason:
                return text, reason
            return text, None

        else:
            return None, "wrong_doc_type"

    except Exception:
        return None, "unreadable"


if __name__ == "__main__":
    from loader import Inbox

    inbox = Inbox("data")

    test_files = [
        ("email_001_SI.txt", False),
        ("email_005_BL.xlsx", False),
        ("email_055_BL.docx", False),
        ("email_059_BL.pdf", False),
        ("email_501_BL.txt", True),  # Commercial Invoice
        ("email_502_BL.txt", True),  # Packing List
        ("email_503_BL.txt", True),  # Certificate of Origin
        ("email_511_BL.pdf", True),  # Corrupted PDF -> unreadable
        ("email_512_BL.pdf", True),  # Image scan -> unreadable
    ]

    print("Running parser smoke tests...")
    for fname, expect_error in test_files:
        path = f"attachments/{fname}"
        b = inbox.read_bytes(path)
        text, reason = parse_attachment(path, b)
        safe_preview = (text or '')[:40].encode('ascii', errors='replace').decode('ascii')
        print(f"[{fname}] size={len(b)} err={reason} preview={repr(safe_preview)}")
        if expect_error:
            assert reason is not None, f"Expected error for {fname}, got None"
        else:
            assert reason is None, f"Expected clean parse for {fname}, got {reason}"
    print("All parser smoke tests passed!")
