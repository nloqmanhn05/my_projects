#!/usr/bin/env python3
"""docio.py — read SI/BL attachments of any supported format into plain text."""
import re

# Binary-format readers are lazily imported so the text path works without deps.


def read_doc(inbox, att_path):
    """Return (raw_text, format, error). Never raises."""
    fmt = att_path.rsplit(".", 1)[-1].lower() if "." in att_path else ""
    try:
        raw = inbox.read_bytes(att_path)
    except FileNotFoundError:
        return "", fmt, "missing_file"
    except Exception as exc:  # pragma: no cover
        return "", fmt, f"read_error: {exc}"
    try:
        if fmt in ("txt", ""):
            return raw.decode("utf-8", errors="replace"), fmt, None
        if fmt == "xlsx":
            return _xlsx_text(raw), fmt, None
        if fmt == "docx":
            return _docx_text(raw), fmt, None
        if fmt in ("pdf",):
            return _pdf_text(raw), fmt, None
        # unknown format: fall back to attempted text decode
        return raw.decode("utf-8", errors="replace"), fmt, None
    except Exception as exc:
        return "", fmt, f"parse_error: {exc}"


def _xlsx_text(raw):
    import io
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    lines = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if not cells:
                continue
            if len(cells) >= 2:
                # first cell is the label, remaining cells form the value
                lines.append(f"{cells[0]} : {chr(10).join(cells[1:])}")
            else:
                lines.append(cells[0])
    return "\n".join(lines)


def _docx_text(raw):
    import io
    import docx
    d = docx.Document(io.BytesIO(raw))
    lines = []
    for p in d.paragraphs:
        if p.text.strip():
            lines.append(p.text)
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            # dedupe merged cells (python-docx repeats cell text for spans)
            uniq = []
            for c in cells:
                if not uniq or uniq[-1] != c:
                    uniq.append(c)
            uniq = [c for c in uniq if c]
            if not uniq:
                continue
            if len(uniq) >= 2:
                lines.append(f"{uniq[0]} : {chr(10).join(uniq[1:])}")
            else:
                lines.append(uniq[0])
    return "\n".join(lines)


def _pdf_text(raw):
    import io
    import pypdf
    reader = pypdf.PdfReader(io.BytesIO(raw))
    pages = []
    for page in reader.pages:
        t = page.extract_text() or ""
        if t.strip():
            pages.append(t)
    text = "\n".join(pages)
    # Scanned PDF (no text layer) — try optional OCR engine if installed.
    if looks_unreadable(text):
        text = _ocr_fallback(raw)
    return text


def _ocr_fallback(raw):
    """Return OCR text for a scanned PDF, or '' when OCR is unavailable."""
    try:
        from ocr import available, ocr_pdf
        if available():
            return ocr_pdf(raw)
    except Exception:
        pass
    return ""


def looks_unreadable(text):
    """True when a doc yields effectively no meaningful text (scanned image, etc.)."""
    stripped = re.sub(r"\s+", " ", text or "").strip()
    return len(stripped) < 30